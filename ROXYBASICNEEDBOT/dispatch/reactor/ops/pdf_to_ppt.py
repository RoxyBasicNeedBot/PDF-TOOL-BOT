# ┌─────────────────────────────────────────────────────────────┐
# │                    𕲏𕲏𕲏𕲏•𝔹𝕒𝕤𝕚𝕔ℕ𝕖𝕖𝕕𝔹𝕠𝕥                     │
# │            The Ultimate Telegram PDF Suite Tool             │
# ├─────────────────────────────────────────────────────────────┤
# │        © 2026 RoxyBasicNeedBot. All Rights Reserved.        │
# └─────────────────────────────────────────────────────────────┘

import os
import fitz
import asyncio
from tracer import tracer

try:
    from pptx import Presentation
    from pptx.util import Inches
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False
    logger.warning("python-pptx not installed. PDF to PowerPoint conversion will not work.")


async def pdfToPPT(input_file: str, cDIR: str) -> (bool, str):
    """
    Convert PDF to PowerPoint by converting each page to a slide image.
    
    This function renders each PDF page as a high-quality image and adds it to a PowerPoint slide.
    Note: Slides contain images (not editable text), suitable for presentation purposes.
    
    Args:
        input_file: Path to input PDF file
        cDIR: User's working directory for temporary files
    
    Returns:
        (True, output_path) on success
        (False, error_message) on failure
    """
    try:
        # Check if library is available
        if not PPTX_AVAILABLE:
            return False, "PDF to PowerPoint converter not available. Please contact bot admin."
        
        # Check if input file exists (with retry for Docker filesystem sync issues)
        for attempt in range(3):
            if os.path.exists(input_file):
                try:
                    with open(input_file, 'rb') as f:
                        f.read(1)  # Read 1 byte to ensure file is accessible
                    break
                except Exception:
                    pass
            await asyncio.sleep(0.5)
        
        if not os.path.exists(input_file):
            logger.error(f"Input file not found after retries: {input_file}")
            return False, "Input file not found. Please try again."
        
        logger.debug(f"📂 Opening PDF for PPT conversion: {input_file}, Size: {os.path.getsize(input_file)}")
        
        # Open PDF
        pdf_document = fitz.open(input_file)
        page_count = len(pdf_document)
        
        output_path = f"{cDIR}/output.pptx"
        
        # Create PowerPoint presentation
        prs = Presentation()
        
        # Set slide dimensions (16:9 aspect ratio - standard for modern presentations)
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(5.625)
        
        # Convert each page to a slide
        for page_num in range(page_count):
            try:
                # Get page
                page = pdf_document[page_num]
                
                # Convert page to high-quality image (2x zoom for better quality)
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                
                # Save temporary image
                img_path = f"{cDIR}/temp_page_{page_num}.png"
                pix.save(img_path)
                
                # Add blank slide
                blank_slide_layout = prs.slide_layouts[6]  # Blank layout
                slide = prs.slides.add_slide(blank_slide_layout)
                
                # Add image to slide (full size)
                left = top = Inches(0)
                slide.shapes.add_picture(
                    img_path,
                    left, top,
                    width=prs.slide_width,
                    height=prs.slide_height
                )
                
                # Clean up temp image
                try:
                    os.remove(img_path)
                except:
                    pass  # Ignore cleanup errors
                    
            except Exception as page_error:
                logger.error(f"Error processing page {page_num + 1}: {page_error}")
                # Continue with next page instead of failing completely
                continue
        
        pdf_document.close()
        
        # Check if any slides were created
        if len(prs.slides) == 0:
            return False, "No slides could be created from PDF"
        
        # Save presentation
        prs.save(output_path)
        
        # Verify output
        if not os.path.exists(output_path) or os.path.getsize(output_path) == 0:
            return False, "Conversion failed: Output file is empty"
        
        return True, output_path
        
    except Exception as e:
        logger.error(f"🐞 Error: {e}", exc_info=True)
        return False, str(Error)
    finally:
        # Clean up any remaining temp images
        try:
            for file in os.listdir(cDIR):
                if file.startswith("temp_page_") and file.endswith(".png"):
                    os.remove(os.path.join(cDIR, file))
        except:
            pass
